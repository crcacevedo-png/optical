from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from bson import ObjectId
from datetime import datetime, timezone
from typing import Optional
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from db import db, serialize_doc
from auth_utils import get_current_user
from cache import TTLCache

router = APIRouter(prefix="/reports", tags=["Reportes"])

# Cache 30s del dashboard, namespaced por empresa+sucursal
dashboard_cache = TTLCache(default_ttl=30, namespace="dashboard")


async def _get_stock_alerts_count(company_id, branch_id=None):
    query = {"company_id": company_id}
    if branch_id:
        query["branch_id"] = branch_id
    stock_items = await db.stock.find(query).to_list(1000)
    if not stock_items:
        return 0
    # Batch fetch de productos para evitar N+1
    pids = list({s["product_id"] for s in stock_items if s.get("product_id")})
    products = await db.products.find(
        {"_id": {"$in": pids}}, {"min_stock": 1}
    ).to_list(len(pids))
    min_stock_map = {p["_id"]: p.get("min_stock", 5) for p in products}
    count = 0
    for s in stock_items:
        if s.get("quantity", 0) <= min_stock_map.get(s.get("product_id"), 5):
            count += 1
    return count


@router.get("/dashboard")
async def get_dashboard(user: dict = Depends(get_current_user), branch_id: Optional[str] = None):
    if user["role"] == "superadmin":
        raise HTTPException(status_code=403, detail="Acceso denegado")

    company_id = ObjectId(user["company_id"])
    effective_branch = branch_id or user.get("branch_id") or "all"
    cache_key = f"{user['company_id']}:{effective_branch}"

    async def _compute():
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        month_start = datetime.now(timezone.utc).strftime("%Y-%m-01")

        branch_filter = {}
        if branch_id:
            try:
                branch_filter["branch_id"] = ObjectId(branch_id)
            except Exception as exc:
                raise HTTPException(status_code=400, detail="branch_id invalido") from exc
        elif user.get("branch_id"):
            branch_filter["branch_id"] = ObjectId(user["branch_id"])

        apt_query = {"company_id": company_id, "date": today, **branch_filter}
        appointments_today = await db.appointments.count_documents(apt_query)
        appointments_pending = await db.appointments.count_documents({**apt_query, "status": "pendiente"})

        patients_query = {"company_id": company_id, "created_at": {"$gte": month_start}}
        new_patients = await db.patients.count_documents(patients_query)

        # Solo sumamos totales — proyeccion + aggregation es mas eficiente que traer docs completos
        sales_today_pipeline = [
            {"$match": {"company_id": company_id, "created_at": {"$gte": today}, **branch_filter}},
            {"$group": {"_id": None, "total": {"$sum": "$total"}, "count": {"$sum": 1}}},
        ]
        today_agg = await db.sales.aggregate(sales_today_pipeline).to_list(1)
        total_sales_today = today_agg[0]["total"] if today_agg else 0
        sales_count_today = today_agg[0]["count"] if today_agg else 0

        sales_month_pipeline = [
            {"$match": {"company_id": company_id, "created_at": {"$gte": month_start}, **branch_filter}},
            {"$group": {"_id": None, "total": {"$sum": "$total"}}},
        ]
        month_agg = await db.sales.aggregate(sales_month_pipeline).to_list(1)
        total_sales_month = month_agg[0]["total"] if month_agg else 0

        finance_pipeline = [
            {"$match": {"company_id": company_id, "date": {"$gte": month_start, "$lte": today}, **branch_filter}},
            {"$group": {"_id": "$type", "total": {"$sum": "$amount"}}},
        ]
        finance_agg = await db.finance_entries.aggregate(finance_pipeline).to_list(10)
        finance_by_type = {r["_id"]: r["total"] for r in finance_agg}
        income = finance_by_type.get("ingreso", 0)
        expense = finance_by_type.get("egreso", 0)

        stock_alerts = await _get_stock_alerts_count(company_id, branch_filter.get("branch_id"))

        upcoming_apts = await db.appointments.find({
            "company_id": company_id, "date": {"$gte": today},
            "status": {"$in": ["pendiente", "confirmada"]},
            **branch_filter
        }).sort([("date", 1), ("time", 1)]).limit(5).to_list(5)

        # Batch fetch pacientes (fix N+1)
        pids = list({a["patient_id"] for a in upcoming_apts if a.get("patient_id")})
        patients_map = {}
        if pids:
            docs = await db.patients.find(
                {"_id": {"$in": pids}}, {"first_name": 1, "last_name": 1}
            ).to_list(len(pids))
            patients_map = {str(p["_id"]): f"{p.get('first_name','')} {p.get('last_name','')}".strip() for p in docs}

        for apt in upcoming_apts:
            serialize_doc(apt)
            pid = apt.get("patient_id")
            if pid:
                apt["patient_name"] = patients_map.get(pid, "")

        return {
            "appointments_today": appointments_today, "appointments_pending": appointments_pending,
            "new_patients": new_patients, "sales_count_today": sales_count_today,
            "total_sales_today": total_sales_today, "total_sales_month": total_sales_month,
            "income": income, "expense": expense, "profit": income - expense,
            "stock_alerts": stock_alerts, "upcoming_appointments": upcoming_apts
        }

    return await dashboard_cache.get_or_load(cache_key, _compute, ttl=30)

@router.get("/sales")
async def get_sales_report(
    user: dict = Depends(get_current_user),
    date_from: str = None,
    date_to: str = None,
    branch_id: Optional[str] = None
):
    if user["role"] == "superadmin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    month_start = datetime.now(timezone.utc).strftime("%Y-%m-01")
    
    query = {"company_id": ObjectId(user["company_id"])}
    if branch_id:
        try:
            query["branch_id"] = ObjectId(branch_id)
        except Exception:
            raise HTTPException(status_code=400, detail="branch_id invalido")
    elif user.get("branch_id"):
        query["branch_id"] = ObjectId(user["branch_id"])
    
    if date_from and date_to:
        query["created_at"] = {"$gte": date_from, "$lte": date_to + "T23:59:59"}
    else:
        query["created_at"] = {"$gte": month_start, "$lte": today + "T23:59:59"}
    
    sales = await db.sales.find(query).to_list(1000)
    total = sum(s.get("total", 0) for s in sales)
    count = len(sales)
    by_payment = {}
    for s in sales:
        pm = s.get("payment_method", "otro")
        by_payment[pm] = by_payment.get(pm, 0) + s.get("total", 0)
    
    return {
        "total": total, "count": count,
        "average": total / count if count > 0 else 0,
        "by_payment_method": by_payment,
        "period": {"from": date_from or month_start, "to": date_to or today}
    }


@router.get("/export/excel")
async def export_reports_excel(
    user: dict = Depends(get_current_user),
    date_from: str = None,
    date_to: str = None,
    branch_id: Optional[str] = None
):
    if user["role"] == "superadmin":
        raise HTTPException(status_code=403, detail="Acceso denegado")
    
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    month_start = datetime.now(timezone.utc).strftime("%Y-%m-01")
    d_from = date_from or month_start
    d_to = date_to or today
    
    company_id = ObjectId(user["company_id"])
    sales_query = {"company_id": company_id}
    finance_query = {"company_id": company_id}
    
    if branch_id:
        try:
            bid = ObjectId(branch_id)
            sales_query["branch_id"] = bid
            finance_query["branch_id"] = bid
        except Exception:
            raise HTTPException(status_code=400, detail="branch_id invalido")
    elif user.get("branch_id"):
        sales_query["branch_id"] = ObjectId(user["branch_id"])
        finance_query["branch_id"] = ObjectId(user["branch_id"])
    
    sales_query["created_at"] = {"$gte": d_from, "$lte": d_to + "T23:59:59"}
    finance_query["date"] = {"$gte": d_from, "$lte": d_to}
    
    sales = await db.sales.find(sales_query).sort("created_at", -1).to_list(5000)
    finance_entries = await db.finance_entries.find(finance_query).sort("date", -1).to_list(5000)
    
    # Resolve patient names
    patient_ids = list({s["patient_id"] for s in sales if s.get("patient_id")})
    patient_map = {}
    if patient_ids:
        patients = await db.patients.find({"_id": {"$in": patient_ids}}, {"first_name": 1, "last_name": 1}).to_list(len(patient_ids))
        patient_map = {p["_id"]: f"{p['first_name']} {p['last_name']}" for p in patients}
    
    # Branch names
    branch_ids = list({s.get("branch_id") for s in sales if s.get("branch_id")} | {e.get("branch_id") for e in finance_entries if e.get("branch_id")})
    branch_map = {}
    if branch_ids:
        branches = await db.branches.find({"_id": {"$in": branch_ids}}, {"name": 1}).to_list(len(branch_ids))
        branch_map = {b["_id"]: b["name"] for b in branches}
    
    # Company name
    company = await db.companies.find_one({"_id": company_id}, {"name": 1})
    company_name = company["name"] if company else "Cortexia Optical"
    
    # Build Excel
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0F4C3A", end_color="0F4C3A", fill_type="solid")
    border = Border(
        bottom=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0')
    )
    
    # -- Sheet 1: Resumen --
    ws1 = wb.active
    ws1.title = "Resumen"
    ws1.append([company_name])
    ws1.merge_cells('A1:D1')
    ws1['A1'].font = Font(bold=True, size=14, color="0F4C3A")
    ws1.append([f"Reporte del {d_from} al {d_to}"])
    ws1['A2'].font = Font(size=10, color="666666")
    ws1.append([])
    
    total_sales = sum(s.get("total", 0) for s in sales)
    income = sum(e["amount"] for e in finance_entries if e["type"] == "ingreso")
    expense = sum(e["amount"] for e in finance_entries if e["type"] == "egreso")
    
    ws1.append(["Metrica", "Valor"])
    for cell in ws1[4]:
        cell.font = header_font
        cell.fill = header_fill
    ws1.append(["Total Ventas (devengado)", f"Q {total_sales:,.2f}"])
    ws1.append(["No. Transacciones", len(sales)])
    ws1.append(["Promedio por Venta", f"Q {(total_sales / len(sales) if sales else 0):,.2f}"])
    ws1.append(["Ingresos cobrados (base caja)", f"Q {income:,.2f}"])
    ws1.append(["Egresos pagados (base caja)", f"Q {expense:,.2f}"])
    ws1.append(["Flujo de caja neto (base caja)", f"Q {income - expense:,.2f}"])
    ws1.append([])
    
    # Payment methods breakdown
    by_payment = {}
    for s in sales:
        pm = s.get("payment_method", "otro")
        label = {"cash": "Efectivo", "card": "Tarjeta", "transfer": "Transferencia"}.get(pm, pm.capitalize())
        by_payment[label] = by_payment.get(label, 0) + s.get("total", 0)
    
    ws1.append(["Ventas por Metodo de Pago", "Monto"])
    row_num = ws1.max_row
    for cell in ws1[row_num]:
        cell.font = header_font
        cell.fill = header_fill
    for method, amount in by_payment.items():
        ws1.append([method, f"Q {amount:,.2f}"])
    
    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 20
    
    # -- Sheet 2: Detalle de Ventas --
    ws2 = wb.create_sheet("Ventas")
    sale_headers = ["Fecha", "Paciente", "Sucursal", "Metodo Pago", "Subtotal", "Descuento", "Total", "Estado"]
    ws2.append(sale_headers)
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    
    for s in sales:
        ws2.append([
            s.get("created_at", "")[:10],
            patient_map.get(s.get("patient_id"), "-"),
            branch_map.get(s.get("branch_id"), "-"),
            {"cash": "Efectivo", "card": "Tarjeta", "transfer": "Transferencia"}.get(s.get("payment_method", ""), s.get("payment_method", "-")),
            round(s.get("subtotal", 0), 2),
            round(s.get("discount", 0), 2),
            round(s.get("total", 0), 2),
            s.get("status", "-")
        ])
    
    for i, w in enumerate([12, 25, 20, 15, 12, 12, 12, 12]):
        ws2.column_dimensions[chr(65 + i)].width = w
    
    # -- Sheet 3: Finanzas --
    ws3 = wb.create_sheet("Finanzas")
    fin_headers = ["Fecha", "Tipo", "Categoria", "Descripcion", "Sucursal", "Monto"]
    ws3.append(fin_headers)
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
    
    for e in finance_entries:
        ws3.append([
            e.get("date", ""),
            "Ingreso" if e["type"] == "ingreso" else "Egreso",
            e.get("category", "-"),
            e.get("description", "-")[:50],
            branch_map.get(e.get("branch_id"), "-"),
            round(e["amount"], 2)
        ])
    
    for i, w in enumerate([12, 10, 15, 35, 20, 12]):
        ws3.column_dimensions[chr(65 + i)].width = w
    
    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"reporte_{company_name.replace(' ', '_')}_{d_from}_{d_to}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

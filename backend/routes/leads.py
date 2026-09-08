"""Captacion de leads (nivel plataforma / global).

- POST /api/leads               -> PUBLICO (sin auth): recibe envios del formulario compartible.
- GET  /api/leads               -> SUPERADMIN: listado con filtros + busqueda.
- GET  /api/leads/grouped       -> SUPERADMIN: envios agrupados por codigo de promocion (+ "sin codigo").
- GET  /api/leads/sources       -> SUPERADMIN: origenes distintos para filtro.
- GET  /api/leads/export        -> SUPERADMIN: exporta a Excel respetando filtros.
- GET/POST/PATCH /api/leads/codes -> SUPERADMIN: gestion de codigos de promocion.

Reglas: correo unico global (409 si repetido), consentimiento obligatorio, honeypot +
rate-limit anti-spam, captura de origen silenciosa desde parametros del enlace.
"""
from fastapi import APIRouter, HTTPException, Depends, Request, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr
from bson import ObjectId
from pymongo.errors import DuplicateKeyError
from datetime import datetime, timezone
from typing import Optional
import io
import re
import os
import asyncio
import secrets
import string
import logging
import html as _html

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from db import db, serialize_doc
from auth_utils import get_current_user, get_real_ip, hash_password
from audit import log_audit
from rate_limiter import limiter
from routes.notifications import create_notification
from email_service import queue_email, render_welcome_company

router = APIRouter(prefix="/leads", tags=["Captacion de Leads"])

logger = logging.getLogger(__name__)

CONSENT_TEXT = (
    "Acepto que Cortexia Optical utilice mis datos para contactarme por WhatsApp y "
    "correo electronico con el fin de gestionar la apertura de mi cuenta. Mis datos "
    "seran tratados de forma confidencial y puedo solicitar su eliminacion cuando lo desee."
)


def _require_superadmin(user: dict):
    if user.get("role") != "superadmin":
        raise HTTPException(status_code=403, detail="Acceso denegado")


def _normalize_wa(raw: str) -> str:
    """Devuelve solo digitos; antepone 502 (Guatemala) si son 8 digitos locales."""
    digits = "".join(ch for ch in (raw or "") if ch.isdigit())
    if len(digits) == 8:
        digits = "502" + digits
    return digits


def _esc(v) -> str:
    return _html.escape(str(v if v is not None else ""))


def _gen_temp_password() -> str:
    """Contrasena temporal fuerte (may/min/digitos + simbolo)."""
    core = "".join(secrets.choice(string.ascii_letters + string.digits) for _ in range(10))
    return f"Cx{core}#7"


async def _notify_superadmins_new_lead(lead: dict, lead_id: str):
    """Encola aviso por correo a los superadmins + notificacion push (best-effort)."""
    recipients = [u["email"] for u in await db.users.find(
        {"role": "superadmin", "is_active": True}, {"email": 1}).to_list(20) if u.get("email")]
    if not recipients:
        fb = os.environ.get("CORTEXIA_ALERTS_TO") or os.environ.get("ADMIN_EMAIL")
        if fb:
            recipients = [fb]
    subject = f"[Cortexia] Nueva solicitud de cuenta: {lead.get('optica_name') or lead.get('name')}"
    html_body = (
        "<h2>Nueva solicitud de cuenta</h2>"
        f"<p><b>Nombre:</b> {_esc(lead.get('name'))}</p>"
        f"<p><b>Optica:</b> {_esc(lead.get('optica_name'))}</p>"
        f"<p><b>Ciudad/Pais:</b> {_esc(lead.get('location'))}</p>"
        f"<p><b>WhatsApp:</b> {_esc(lead.get('whatsapp'))}</p>"
        f"<p><b>Correo:</b> {_esc(lead.get('email'))}</p>"
        f"<p><b>Codigo:</b> {_esc(lead.get('promo_code') or 'Sin codigo')}</p>"
        f"<p><b>Origen:</b> {_esc(lead.get('source'))}</p>"
        "<p>Revisala en el panel de SuperAdmin &gt; Solicitudes.</p>"
    )
    for em in recipients:
        await queue_email(em, subject, html_body, tag="new_lead")
    await create_notification(
        "new_lead", "Nueva solicitud de cuenta",
        f"{lead.get('name')} ({lead.get('optica_name')}) solicito abrir cuenta",
        {"lead_id": lead_id, "email": lead.get("email")},
    )


# ─────────────────────────── FORMULARIO PUBLICO ───────────────────────────
class LeadSubmit(BaseModel):
    name: str
    optica_name: Optional[str] = ""
    location: Optional[str] = ""
    whatsapp: str
    email: EmailStr
    promo_code: Optional[str] = ""
    consent: bool = False
    source: Optional[str] = None
    source_details: Optional[dict] = None
    website: Optional[str] = ""  # honeypot: los humanos lo dejan vacio


@router.post("")
@limiter.limit("5/minute")
async def submit_lead(request: Request, data: LeadSubmit):
    # Honeypot: si viene lleno es un bot => fingimos exito sin guardar nada.
    if data.website:
        return {"ok": True, "message": "Gracias por registrarte."}

    name = (data.name or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="El nombre es obligatorio.")
    optica_name = (data.optica_name or "").strip()
    if not optica_name:
        raise HTTPException(status_code=400, detail="El nombre de la optica es obligatorio.")
    location = (data.location or "").strip()
    if not location:
        raise HTTPException(status_code=400, detail="La ciudad / pais es obligatoria.")
    if not data.consent:
        raise HTTPException(status_code=400, detail="Debes aceptar el consentimiento para continuar.")

    wa_digits = _normalize_wa(data.whatsapp)
    if len(wa_digits) < 8:
        raise HTTPException(status_code=400, detail="El numero de WhatsApp no es valido.")
    whatsapp = "+" + wa_digits

    email = str(data.email).strip().lower()
    if await db.leads.find_one({"email": email}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Este correo ya esta registrado.")

    # Posible duplicado por WhatsApp (se permite, pero se marca).
    is_dup = bool(await db.leads.find_one({"whatsapp": whatsapp}, {"_id": 1}))

    # Validar codigo escrito contra codigos ACTIVOS predefinidos.
    code_input = (data.promo_code or "").strip().upper()
    matched = None
    if code_input:
        pc = await db.promo_codes.find_one({"code": code_input, "is_active": True}, {"code": 1})
        if pc:
            matched = pc["code"]

    source = (data.source or "").strip() or "no_especificado"
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "name": name,
        "optica_name": optica_name,
        "location": location,
        "whatsapp": whatsapp,
        "whatsapp_raw": (data.whatsapp or "").strip(),
        "email": email,
        "promo_code": matched,
        "promo_code_input": code_input or None,
        "consent": True,
        "consent_text": CONSENT_TEXT,
        "source": source,
        "source_details": data.source_details or {},
        "is_possible_duplicate": is_dup,
        "status": "nueva",
        "ip": get_real_ip(request),
        "user_agent": (request.headers.get("user-agent") or "")[:300],
        "created_at": now,
    }
    try:
        res = await db.leads.insert_one(doc)
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail="Este correo ya esta registrado.")

    try:
        await _notify_superadmins_new_lead(doc, str(res.inserted_id))
    except Exception as e:
        logger.warning(f"No se pudo notificar nueva solicitud: {e}")

    return {"ok": True, "message": "Gracias por registrarte. Pronto nos pondremos en contacto contigo."}


# ─────────────────────────── ADMIN (SUPERADMIN) ───────────────────────────
def _build_leads_query(promo_code, source, search, duplicates, status=None) -> dict:
    q: dict = {}
    if promo_code:
        if promo_code in ("none", "sin_codigo", "null"):
            q["promo_code"] = None
        else:
            q["promo_code"] = promo_code.upper()
    if source:
        q["source"] = source
    if status and status != "all":
        q["status"] = status
    if duplicates:
        q["is_possible_duplicate"] = True
    if search:
        rx = {"$regex": re.escape(search), "$options": "i"}
        q["$or"] = [{"name": rx}, {"email": rx}, {"whatsapp": rx}, {"whatsapp_raw": rx}, {"optica_name": rx}]
    return q


@router.get("")
async def list_leads(
    user: dict = Depends(get_current_user),
    promo_code: Optional[str] = Query(None),
    source: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    duplicates: Optional[bool] = Query(None),
    limit: int = Query(100, ge=1, le=500),
    skip: int = Query(0, ge=0),
):
    _require_superadmin(user)
    q = _build_leads_query(promo_code, source, search, duplicates, status)
    total = await db.leads.count_documents(q)
    docs = await db.leads.find(q).sort("created_at", -1).skip(skip).limit(limit).to_list(limit)
    return {"items": [serialize_doc(d) for d in docs], "total": total, "skip": skip, "limit": limit}


@router.get("/stats")
async def leads_stats(user: dict = Depends(get_current_user)):
    _require_superadmin(user)
    total = await db.leads.count_documents({})
    with_code = await db.leads.count_documents({"promo_code": {"$ne": None}})
    dups = await db.leads.count_documents({"is_possible_duplicate": True})
    return {"total": total, "with_code": with_code, "without_code": total - with_code, "possible_duplicates": dups}


def _member(m: dict) -> dict:
    return {
        "id": str(m["_id"]),
        "name": m.get("name"),
        "optica_name": m.get("optica_name"),
        "location": m.get("location"),
        "whatsapp": m.get("whatsapp"),
        "email": m.get("email"),
        "source": m.get("source"),
        "created_at": m.get("created_at"),
        "status": m.get("status", "nueva"),
        "is_possible_duplicate": m.get("is_possible_duplicate", False),
    }


@router.get("/grouped")
async def grouped_leads(user: dict = Depends(get_current_user)):
    _require_superadmin(user)
    codes = await db.promo_codes.find().sort("code", 1).to_list(None)
    groups = []
    for pc in codes:
        members = await db.leads.find({"promo_code": pc["code"]}).sort("created_at", -1).to_list(None)
        groups.append({
            "code": pc["code"],
            "label": pc.get("label", ""),
            "is_active": pc.get("is_active", True),
            "count": len(members),
            "members": [_member(m) for m in members],
        })
    no_code = await db.leads.find({"promo_code": None}).sort("created_at", -1).to_list(None)
    groups.append({
        "code": None, "label": "Sin codigo", "is_active": True,
        "count": len(no_code), "members": [_member(m) for m in no_code],
    })
    return {"groups": groups}


@router.get("/sources")
async def lead_sources(user: dict = Depends(get_current_user)):
    _require_superadmin(user)
    vals = await db.leads.distinct("source")
    return {"sources": sorted([v for v in vals if v])}


def _render_leads_xlsx(docs: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    headers = ["Nombre", "Optica", "Ciudad/Pais", "WhatsApp", "Correo", "Codigo", "Origen", "Estado", "Posible duplicado", "Consentimiento", "Fecha"]
    fill = PatternFill(start_color="1B2A49", end_color="1B2A49", fill_type="solid")
    hfont = Font(bold=True, color="FFFFFF")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fill
        c.font = hfont
        c.alignment = Alignment(horizontal="center")
    _status_lbl = {"nueva": "Nueva", "contactada": "Contactada", "cuenta_creada": "Cuenta creada"}
    for r, d in enumerate(docs, start=2):
        ws.cell(row=r, column=1, value=d.get("name"))
        ws.cell(row=r, column=2, value=d.get("optica_name"))
        ws.cell(row=r, column=3, value=d.get("location"))
        ws.cell(row=r, column=4, value=d.get("whatsapp"))
        ws.cell(row=r, column=5, value=d.get("email"))
        ws.cell(row=r, column=6, value=d.get("promo_code") or "Sin codigo")
        ws.cell(row=r, column=7, value=d.get("source") or "no_especificado")
        ws.cell(row=r, column=8, value=_status_lbl.get(d.get("status", "nueva"), "Nueva"))
        ws.cell(row=r, column=9, value="Si" if d.get("is_possible_duplicate") else "No")
        ws.cell(row=r, column=10, value="Si" if d.get("consent") else "No")
        ws.cell(row=r, column=11, value=(d.get("created_at") or "")[:19].replace("T", " "))
    widths = [24, 24, 22, 16, 30, 14, 18, 14, 16, 14, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@router.get("/export")
async def export_leads(
    request: Request,
    user: dict = Depends(get_current_user),
    promo_code: Optional[str] = Query(None),
    source: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    duplicates: Optional[bool] = Query(None),
):
    _require_superadmin(user)
    q = _build_leads_query(promo_code, source, search, duplicates, status)
    docs = await db.leads.find(q).sort("created_at", -1).to_list(None)
    xlsx = await asyncio.to_thread(_render_leads_xlsx, docs)
    await log_audit("LEADS_EXPORTED", actor_id=user["_id"], actor_email=user.get("email"),
                    actor_role="superadmin", metadata={"count": len(docs), "filters": q}, request=request)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(
        io.BytesIO(xlsx),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Cortexia_Leads_{ts}.xlsx"'},
    )


# ─────────────────────────── CODIGOS DE PROMOCION ───────────────────────────
class PromoCodeCreate(BaseModel):
    code: str
    label: Optional[str] = ""


class PromoCodeUpdate(BaseModel):
    is_active: Optional[bool] = None
    label: Optional[str] = None


@router.get("/codes")
async def list_promo_codes(user: dict = Depends(get_current_user)):
    _require_superadmin(user)
    codes = await db.promo_codes.find().sort("created_at", -1).to_list(None)
    for c in codes:
        c["lead_count"] = await db.leads.count_documents({"promo_code": c["code"]})
    return {"items": [serialize_doc(c) for c in codes]}


@router.post("/codes")
async def create_promo_code(request: Request, data: PromoCodeCreate, user: dict = Depends(get_current_user)):
    _require_superadmin(user)
    code = (data.code or "").strip().upper()
    if not code:
        raise HTTPException(status_code=400, detail="El codigo es obligatorio.")
    if not re.match(r"^[A-Z0-9_-]{2,32}$", code):
        raise HTTPException(status_code=400, detail="El codigo solo admite letras, numeros, guion y guion bajo (2-32).")
    if await db.promo_codes.find_one({"code": code}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Ese codigo ya existe.")
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "code": code,
        "label": (data.label or "").strip(),
        "is_active": True,
        "created_at": now,
        "created_by": user["_id"],
        "updated_at": now,
    }
    try:
        res = await db.promo_codes.insert_one(doc)
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail="Ese codigo ya existe.")
    await log_audit("PROMO_CODE_CREATED", actor_id=user["_id"], actor_email=user.get("email"),
                    actor_role="superadmin", metadata={"code": code}, request=request)
    doc["_id"] = res.inserted_id
    doc["lead_count"] = 0
    return serialize_doc(doc)


@router.patch("/codes/{code_id}")
async def update_promo_code(code_id: str, data: PromoCodeUpdate, request: Request, user: dict = Depends(get_current_user)):
    _require_superadmin(user)
    try:
        oid = ObjectId(code_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID invalido")
    upd = {"updated_at": datetime.now(timezone.utc).isoformat()}
    if data.is_active is not None:
        upd["is_active"] = data.is_active
    if data.label is not None:
        upd["label"] = data.label.strip()
    r = await db.promo_codes.update_one({"_id": oid}, {"$set": upd})
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Codigo no encontrado")
    await log_audit("PROMO_CODE_UPDATED", actor_id=user["_id"], actor_email=user.get("email"),
                    actor_role="superadmin", metadata={"code_id": code_id, "changes": {k: v for k, v in upd.items() if k != "updated_at"}}, request=request)
    return {"ok": True}


# ─────────────────────────── ESTADO Y ALTA DE CUENTA ───────────────────────────
class StatusUpdate(BaseModel):
    status: str


class CreateAccountRequest(BaseModel):
    admin_password: Optional[str] = None


_VALID_STATUS = {"nueva", "contactada", "cuenta_creada"}


@router.patch("/{lead_id}/status")
async def update_lead_status(lead_id: str, data: StatusUpdate, request: Request, user: dict = Depends(get_current_user)):
    _require_superadmin(user)
    if data.status not in _VALID_STATUS:
        raise HTTPException(status_code=400, detail="Estado invalido.")
    try:
        oid = ObjectId(lead_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID invalido")
    r = await db.leads.update_one(
        {"_id": oid},
        {"$set": {"status": data.status, "status_updated_at": datetime.now(timezone.utc).isoformat()}},
    )
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    await log_audit("LEAD_STATUS_UPDATED", actor_id=user["_id"], actor_email=user.get("email"),
                    actor_role="superadmin", metadata={"lead_id": lead_id, "status": data.status}, request=request)
    return {"ok": True}


@router.post("/{lead_id}/create-account")
async def create_account_from_lead(lead_id: str, data: CreateAccountRequest, request: Request, user: dict = Depends(get_current_user)):
    """Crea la optica (empresa) + usuario admin a partir de una solicitud.
    Reutiliza el mismo flujo de alta que POST /api/companies (hash_password + correo de bienvenida)."""
    _require_superadmin(user)
    try:
        oid = ObjectId(lead_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID invalido")
    lead = await db.leads.find_one({"_id": oid})
    if not lead:
        raise HTTPException(status_code=404, detail="Solicitud no encontrada")
    if lead.get("company_id"):
        raise HTTPException(status_code=409, detail="Esta solicitud ya tiene una cuenta creada.")

    admin_email = (lead.get("email") or "").strip().lower()
    if not admin_email:
        raise HTTPException(status_code=400, detail="La solicitud no tiene correo.")
    if await db.users.find_one({"email": admin_email}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Ya existe un usuario con ese correo; no se puede crear la cuenta automaticamente.")

    company_name = (lead.get("optica_name") or lead.get("name") or "Optica").strip()
    admin_name = (lead.get("name") or "Administrador").strip()
    phone = lead.get("whatsapp") or ""
    password = (data.admin_password or "").strip() or _gen_temp_password()
    now = datetime.now(timezone.utc).isoformat()

    company_doc = {
        "name": company_name, "legal_name": "", "tax_id": "", "address": lead.get("location") or "",
        "phone": phone, "email": admin_email, "contact_name": admin_name,
        "contact_phone": phone, "contact_email": admin_email,
        "is_active": True, "created_at": now,
    }
    result = await db.companies.insert_one(company_doc)
    company_id = result.inserted_id

    admin_doc = {
        "email": admin_email, "password_hash": hash_password(password), "name": admin_name,
        "role": "admin", "company_id": company_id, "branch_id": None,
        "is_active": True, "created_at": now,
    }
    try:
        await db.users.insert_one(admin_doc)
    except Exception:
        await db.companies.delete_one({"_id": company_id})  # rollback si el correo colisiona
        raise HTTPException(status_code=409, detail="Ya existe un usuario con ese correo; no se pudo crear la cuenta.")

    await create_notification(
        "new_company", "Nueva optica creada desde solicitud",
        f"{company_name} - admin {admin_name} ({admin_email})",
        {"company_id": str(company_id), "company_name": company_name},
    )
    app_url = os.environ.get("APP_URL", "https://cortexiaoptical.com")
    welcome_html = render_welcome_company(
        admin_name=admin_name, company_name=company_name,
        admin_email=admin_email, admin_password=password, login_link=app_url,
    )
    await queue_email(admin_email, f"Bienvenido a Cortexia Optical - {company_name}", welcome_html, tag="welcome")

    await db.leads.update_one({"_id": oid}, {"$set": {
        "status": "cuenta_creada", "company_id": company_id,
        "converted_at": now, "converted_by": user["_id"], "status_updated_at": now,
    }})
    await log_audit("LEAD_ACCOUNT_CREATED", actor_id=user["_id"], actor_email=user.get("email"),
                    actor_role="superadmin",
                    metadata={"lead_id": lead_id, "company_id": str(company_id), "admin_email": admin_email},
                    request=request)
    return {
        "ok": True, "company_id": str(company_id), "admin_email": admin_email,
        "temp_password": password,
        "message": "Cuenta creada. Se envio el correo de bienvenida con las credenciales.",
    }

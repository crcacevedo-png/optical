"""Backend tests for Excel export and plan history (iteration 14)."""
import os
import requests
import pytest

def _load_base_url():
    url = os.environ.get('REACT_APP_BACKEND_URL', '').strip()
    if not url:
        try:
            with open('/app/frontend/.env') as f:
                for line in f:
                    if line.startswith('REACT_APP_BACKEND_URL='):
                        url = line.split('=', 1)[1].strip()
                        break
        except Exception:
            pass
    assert url, "REACT_APP_BACKEND_URL not set"
    return url.rstrip('/')


BASE_URL = _load_base_url()
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, f"Login failed for {email}: {r.text}"
    return s, r.json()


@pytest.fixture(scope="module")
def superadmin():
    s, u = login("superadmin@cortexia.com", "Admin123!")
    return s, u


@pytest.fixture(scope="module")
def altavista_admin():
    s, u = login("analuhs@gmail.com", "Alta2026$")
    return s, u


@pytest.fixture(scope="module")
def demo_admin():
    s, u = login("admin@cortexia.gt", "Demo123!")
    return s, u


# ---------- Excel Export ----------
class TestExcelExport:
    def test_export_excel_returns_xlsx(self, altavista_admin):
        s, _ = altavista_admin
        r = s.get(f"{BASE_URL}/api/reports/export/excel")
        assert r.status_code == 200, r.text
        assert r.headers.get("content-type", "").startswith(EXCEL_MIME)
        # PK signature for xlsx (zip)
        assert r.content[:2] == b"PK", "Response is not a valid xlsx"
        assert "attachment" in r.headers.get("content-disposition", "").lower()
        assert ".xlsx" in r.headers.get("content-disposition", "")

    def test_export_excel_with_date_range(self, altavista_admin):
        s, _ = altavista_admin
        r = s.get(f"{BASE_URL}/api/reports/export/excel",
                  params={"date_from": "2025-01-01", "date_to": "2026-12-31"})
        assert r.status_code == 200
        assert r.content[:2] == b"PK"
        # filename should include dates
        cd = r.headers.get("content-disposition", "")
        assert "2025-01-01" in cd and "2026-12-31" in cd

    def test_export_excel_three_sheets(self, altavista_admin):
        from openpyxl import load_workbook
        import io
        s, _ = altavista_admin
        r = s.get(f"{BASE_URL}/api/reports/export/excel")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert set(["Resumen", "Ventas", "Finanzas"]).issubset(set(wb.sheetnames)), \
            f"Expected sheets missing. Got {wb.sheetnames}"

    def test_export_excel_branch_filter_invalid(self, altavista_admin):
        s, _ = altavista_admin
        r = s.get(f"{BASE_URL}/api/reports/export/excel",
                  params={"branch_id": "not-a-valid-id"})
        assert r.status_code == 400

    def test_export_excel_branch_filter_valid(self, altavista_admin):
        s, _ = altavista_admin
        br = s.get(f"{BASE_URL}/api/branches")
        assert br.status_code == 200
        branches = br.json()
        if not branches:
            pytest.skip("No branches available")
        bid = branches[0]["_id"]
        r = s.get(f"{BASE_URL}/api/reports/export/excel", params={"branch_id": bid})
        assert r.status_code == 200
        assert r.content[:2] == b"PK"

    def test_export_excel_superadmin_forbidden(self, superadmin):
        s, _ = superadmin
        r = s.get(f"{BASE_URL}/api/reports/export/excel")
        assert r.status_code == 403

    def test_export_excel_unauthenticated(self):
        r = requests.get(f"{BASE_URL}/api/reports/export/excel")
        assert r.status_code in (401, 403)


# ---------- Plan History ----------
class TestPlanHistory:
    def test_get_history_superadmin(self, superadmin):
        s, _ = superadmin
        comp = s.get(f"{BASE_URL}/api/companies").json()
        altavista = next((c for c in comp if "altavista" in c.get("name", "").lower()), None)
        assert altavista, "Altavista company not found"
        r = s.get(f"{BASE_URL}/api/plans/history/{altavista['_id']}")
        assert r.status_code == 200
        history = r.json()
        assert isinstance(history, list)
        if len(history) >= 2:
            # should be sorted desc by changed_at
            assert history[0]["changed_at"] >= history[1]["changed_at"]
            entry = history[0]
            for k in ("new_plan_name", "changed_at", "company_id"):
                assert k in entry, f"Missing field {k} in history entry"
            # _id should be string (not ObjectId)
            assert isinstance(entry.get("_id"), str)

    def test_history_forbidden_for_admin(self, altavista_admin):
        s, u = altavista_admin
        cid = u.get("user", {}).get("company_id") or u.get("company_id")
        if not cid:
            comp = requests.get(f"{BASE_URL}/api/companies")
            pytest.skip("No company_id found")
        r = s.get(f"{BASE_URL}/api/plans/history/{cid}")
        assert r.status_code == 403

    def test_history_invalid_company_id(self, superadmin):
        s, _ = superadmin
        r = s.get(f"{BASE_URL}/api/plans/history/not-an-objectid")
        assert r.status_code == 400

    def test_assign_plan_creates_history_entry(self, superadmin):
        s, _ = superadmin
        comp = s.get(f"{BASE_URL}/api/companies").json()
        altavista = next((c for c in comp if "altavista" in c.get("name", "").lower()), None)
        assert altavista
        cid = altavista["_id"]

        plans = s.get(f"{BASE_URL}/api/plans").json()
        free = next(p for p in plans if p["name"].lower() == "free")
        enterprise = next(p for p in plans if p["name"].lower() == "enterprise")

        before = s.get(f"{BASE_URL}/api/plans/history/{cid}").json()
        before_count = len(before)

        # assign Free
        r1 = s.put(f"{BASE_URL}/api/plans/assign/{cid}", params={"plan_id": free["_id"]})
        assert r1.status_code == 200, r1.text
        # restore Enterprise
        r2 = s.put(f"{BASE_URL}/api/plans/assign/{cid}", params={"plan_id": enterprise["_id"]})
        assert r2.status_code == 200, r2.text

        after = s.get(f"{BASE_URL}/api/plans/history/{cid}").json()
        assert len(after) == before_count + 2

        latest = after[0]
        assert latest["new_plan_name"] == enterprise["name"]
        assert latest.get("old_plan_name") in (free["name"], None, "Free")
        assert "changed_by_name" in latest
        assert "changed_at" in latest

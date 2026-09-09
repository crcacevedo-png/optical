"""Iter13: GET /api/finance/expenses-report.xlsx (rango de fechas, cierre contable)
- admin/vendedor: 200 xlsx (PK) con hojas 'Resumen' + 'Detalle'
- excluye ingresos, excluye anulados, respeta branch_id
- superadmin -> 403; branch_id invalido -> 400
"""
import io
import os
import pytest
import requests
from datetime import datetime, timezone

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ADMIN = ("admin@cortexia.gt", "DemoAdmin2026!")
VEND = ("vendedor@cortexia.gt", "DemoUser2026!")
SUPER = ("superadmin@cortexia.com", "Montecristo2026")


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text[:200]}"
    csrf = s.cookies.get("csrf_token")
    if csrf:
        s.headers.update({"X-CSRF-Token": csrf})
    return s


def _parse(content):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    data = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([("" if c is None else c) for c in row])
        data[ws.title] = rows
    return data


_state = {"ids": [], "voided": None,
          "descs": ["QA13 alfa cash", "QA13 beta transfer", "QA13 gamma anulado", "QA13 delta INGRESO"]}


@pytest.fixture(scope="module")
def admin():
    return _login(*ADMIN)


@pytest.fixture(scope="module")
def vendedor():
    return _login(*VEND)


@pytest.fixture(scope="module")
def superadmin():
    return _login(*SUPER)


class TestSeed:
    def test_01_seed(self, admin):
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        payloads = [
            {"type": "egreso", "category": "other_expense", "amount": 30, "description": _state["descs"][0], "date": today, "payment_method": "cash"},
            {"type": "egreso", "category": "other_expense", "amount": 20, "description": _state["descs"][1], "date": today, "payment_method": "transfer"},
            {"type": "egreso", "category": "other_expense", "amount": 10, "description": _state["descs"][2], "date": today, "payment_method": "cash"},
            {"type": "ingreso", "category": "sale", "amount": 999, "description": _state["descs"][3], "date": today, "payment_method": "cash"},
        ]
        for p in payloads:
            r = admin.post(f"{BASE_URL}/api/finance", json=p, timeout=30)
            assert r.status_code == 200, r.text[:200]
            _state["ids"].append(r.json()["_id"])
        # anular el 3ro
        vid = _state["ids"][2]
        _state["voided"] = vid
        r = admin.post(f"{BASE_URL}/api/finance/{vid}/void", json={"reason": "QA13"}, timeout=30)
        assert r.status_code == 200


class TestReport:
    def test_02_admin_xlsx(self, admin):
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        r = admin.get(f"{BASE_URL}/api/finance/expenses-report.xlsx",
                      params={"date_from": today, "date_to": today}, timeout=60)
        assert r.status_code == 200, r.text[:200]
        assert r.content[:2] == b"PK"
        assert XLSX_CT in r.headers.get("content-type", "").lower()
        data = _parse(r.content)
        assert "Resumen" in data and "Detalle" in data, list(data.keys())
        resumen_flat = "\n".join(" | ".join(str(x) for x in row) for row in data["Resumen"])
        detalle_flat = "\n".join(" | ".join(str(x) for x in row) for row in data["Detalle"])
        # Resumen contains TOTAL EGRESOS, Por Categoria, Por Metodo de Pago
        assert "TOTAL EGRESOS" in resumen_flat
        assert "Por Categoria" in resumen_flat
        assert "Por Metodo de Pago" in resumen_flat
        # Detalle contains alfa & beta but NOT anulado gamma nor ingreso delta
        assert "QA13 alfa cash" in detalle_flat
        assert "QA13 beta transfer" in detalle_flat
        assert "QA13 gamma anulado" not in detalle_flat, "anulado presente!"
        assert "QA13 delta INGRESO" not in detalle_flat, "ingreso presente!"
        # header row present
        assert data["Detalle"][0][:5] == ["Fecha", "Categoria", "Proveedor", "Descripcion", "Metodo"]
        assert data["Detalle"][0][5:10] == ["Credito", "Pagado", "Saldo", "Referencia", "Monto"]

    def test_03_vendedor_xlsx(self, vendedor):
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        r = vendedor.get(f"{BASE_URL}/api/finance/expenses-report.xlsx",
                         params={"date_from": today, "date_to": today}, timeout=60)
        assert r.status_code == 200
        assert r.content[:2] == b"PK"

    def test_04_superadmin_403(self, superadmin):
        r = superadmin.get(f"{BASE_URL}/api/finance/expenses-report.xlsx", timeout=30)
        assert r.status_code == 403

    def test_05_branch_invalid_400(self, admin):
        r = admin.get(f"{BASE_URL}/api/finance/expenses-report.xlsx",
                      params={"branch_id": "not-an-oid"}, timeout=30)
        assert r.status_code == 400

    def test_06_regression_purchases_xlsx(self, admin):
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        r = admin.get(f"{BASE_URL}/api/finance/purchases-report.xlsx",
                      params={"date_from": today, "date_to": today}, timeout=60)
        assert r.status_code == 200, r.text[:200]
        assert r.content[:2] == b"PK"
        assert XLSX_CT in r.headers.get("content-type", "").lower()


class TestZZCleanup:
    def test_zz(self):
        try:
            import asyncio
            from motor.motor_asyncio import AsyncIOMotorClient
            from bson import ObjectId as _OID
            mongo_url = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
            db_name = os.environ.get("DB_NAME", "opticas_saas")

            async def _clean():
                client = AsyncIOMotorClient(mongo_url)
                db = client[db_name]
                eids = [_OID(x) for x in _state["ids"]]
                if eids:
                    await db.finance_entries.delete_many({"_id": {"$in": eids}})
                # limpieza extra defensiva por prefijo QA13
                await db.finance_entries.delete_many({"description": {"$regex": "^QA13"}})
                client.close()

            try:
                loop = asyncio.get_event_loop()
                if loop.is_closed():
                    raise RuntimeError("closed")
            except RuntimeError:
                loop = asyncio.new_event_loop()
                asyncio.set_event_loop(loop)
            loop.run_until_complete(_clean())
        except Exception as ex:
            print(f"cleanup err: {ex}")

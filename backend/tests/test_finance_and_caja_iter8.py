"""Regression tests iter8: Finance (A-H) + Cash Register (1-3) observational.

Escenarios:
 - FINANZAS A: proveedor obligatorio en egresos categoría 'suppliers'
 - FINANZAS E: egreso a crédito (devengado) reflejado en summary
 - FINANZAS F: /finance/payables listing
 - FINANZAS G: abono a cuenta por pagar
 - FINANZAS H: /finance/payables/alerts
 - FINANZAS D: /finance/purchases-report.xlsx
 - CAJA 1-3: apertura/cierre + observación egresos NO reflejados
"""
import os
import io
import pytest
import requests
from datetime import datetime, timezone, timedelta

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")

VENDEDOR_EMAIL = os.getenv("TEST_VENDEDOR_EMAIL", "vendedor@cortexia.gt")
VENDEDOR_PASSWORD = os.getenv("TEST_VENDEDOR_PASSWORD") or os.getenv("DEMO_PASSWORD") or "DemoUser2026!"
ADMIN_EMAIL = os.getenv("TEST_ADMIN_EMAIL", "admin@cortexia.gt")
ADMIN_PASSWORD = os.getenv("TEST_ADMIN_PASSWORD") or os.getenv("DEMO_PASSWORD") or "DemoAdmin2026!"


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed {email}: {r.status_code} {r.text[:200]}"
    # attach csrf header for subsequent mutating requests
    csrf = s.cookies.get("csrf_token")
    if csrf:
        s.headers.update({"X-CSRF-Token": csrf})
    return s, r.json()


def _ensure_closed(sess):
    try:
        r = sess.get(f"{BASE_URL}/api/cash-register/current", timeout=30)
        if r.status_code == 200 and r.json().get("register"):
            sess.post(f"{BASE_URL}/api/cash-register/close",
                      json={"counted_cash": 0, "closing_notes": "cleanup iter8"}, timeout=30)
    except Exception:
        pass


@pytest.fixture(scope="module")
def vendedor():
    s, user = _login(VENDEDOR_EMAIL, VENDEDOR_PASSWORD)
    _ensure_closed(s)
    yield s, user
    _ensure_closed(s)


@pytest.fixture(scope="module")
def admin():
    s, user = _login(ADMIN_EMAIL, ADMIN_PASSWORD)
    yield s, user


@pytest.fixture(scope="module")
def qa_supplier(admin):
    s, _ = admin
    r = s.post(f"{BASE_URL}/api/suppliers", json={"name": "QA Proveedor Iter8", "phone": "5555-0000"}, timeout=30)
    assert r.status_code in (200, 201), f"create supplier: {r.status_code} {r.text[:200]}"
    data = r.json()
    sid = data.get("_id") or data.get("id") or data.get("supplier_id")
    assert sid, f"no supplier id in {data}"
    yield sid
    # cleanup
    try:
        s.delete(f"{BASE_URL}/api/suppliers/{sid}", timeout=30)
    except Exception:
        pass


def _cleanup_admin(admin_sess):
    """Delete admin cleanup helper."""
    pass


# ================= FINANZAS A =================
class TestFinanzasA:
    def test_egreso_suppliers_sin_supplier_id_400(self, vendedor):
        s, _ = vendedor
        payload = {"type": "egreso", "category": "suppliers", "amount": 100,
                   "description": "QA sin proveedor"}
        r = s.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r.status_code == 400, f"expected 400 got {r.status_code}: {r.text[:200]}"
        assert "proveedor" in r.text.lower()

    def test_egreso_suppliers_id_invalido_400(self, vendedor):
        s, _ = vendedor
        payload = {"type": "egreso", "category": "suppliers", "amount": 100,
                   "description": "QA inv", "supplier_id": "not-an-oid"}
        r = s.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r.status_code == 400, f"got {r.status_code}: {r.text[:200]}"

    def test_egreso_suppliers_id_inexistente_404(self, vendedor):
        s, _ = vendedor
        payload = {"type": "egreso", "category": "suppliers", "amount": 100,
                   "description": "QA inex", "supplier_id": "507f1f77bcf86cd799439011"}
        r = s.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r.status_code == 404, f"got {r.status_code}: {r.text[:200]}"

    def test_egreso_suppliers_ok_incluye_supplier_name(self, vendedor, qa_supplier):
        s, _ = vendedor
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        payload = {"type": "egreso", "category": "suppliers", "amount": 250,
                   "description": "QA compra ok", "supplier_id": qa_supplier, "date": today}
        r = s.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r.status_code == 200, f"got {r.status_code}: {r.text[:200]}"
        eid = r.json().get("_id")
        assert eid
        # verify persisted
        r2 = s.get(f"{BASE_URL}/api/finance", params={"type": "egreso"}, timeout=30)
        assert r2.status_code == 200
        found = [e for e in r2.json() if e.get("_id") == eid]
        assert found, "created egreso not returned in list"
        assert found[0].get("supplier_name") == "QA Proveedor Iter8"
        # cleanup
        try:
            s.delete(f"{BASE_URL}/api/finance/{eid}", timeout=30)
        except Exception:
            pass


# ================= FINANZAS E, F, G, H =================
class TestPayables:
    _entry_id = None
    _amount_total = 1000
    _amount_paid = 200

    def test_e_crear_egreso_credito(self, vendedor, qa_supplier):
        s, _ = vendedor
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        due = (datetime.now(timezone.utc) + timedelta(days=3)).strftime("%Y-%m-%d")
        # summary antes
        r0 = s.get(f"{BASE_URL}/api/finance/summary", timeout=30)
        assert r0.status_code == 200
        expense_before = float(r0.json().get("expense", 0))

        payload = {"type": "egreso", "category": "suppliers",
                   "amount": self._amount_total, "description": "QA credito iter8",
                   "supplier_id": qa_supplier, "date": today,
                   "is_credit": True, "amount_paid": self._amount_paid, "due_date": due}
        r = s.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        eid = r.json()["_id"]
        TestPayables._entry_id = eid

        # summary aumenta el TOTAL (1000) inmediato (devengado)
        r1 = s.get(f"{BASE_URL}/api/finance/summary", timeout=30)
        assert r1.status_code == 200
        expense_after = float(r1.json().get("expense", 0))
        assert expense_after - expense_before >= self._amount_total - 0.01, \
            f"devengado: expected +{self._amount_total}, got +{expense_after-expense_before}"

    def test_f_payables_list(self, vendedor):
        s, _ = vendedor
        r = s.get(f"{BASE_URL}/api/finance/payables", timeout=30)
        assert r.status_code == 200, r.text[:200]
        body = r.json()
        assert "items" in body and "total_pending" in body and "count" in body and "by_supplier" in body
        eid = TestPayables._entry_id
        found = [x for x in body["items"] if x.get("_id") == eid]
        assert found, f"payable {eid} not in list"
        assert abs(float(found[0]["balance"]) - (self._amount_total - self._amount_paid)) < 0.01

    def test_h_alerts(self, vendedor):
        s, _ = vendedor
        r = s.get(f"{BASE_URL}/api/finance/payables/alerts", params={"days": 7}, timeout=30)
        assert r.status_code == 200, r.text[:200]
        body = r.json()
        assert body["days"] == 7
        assert "overdue" in body and "due_soon" in body
        # el egreso creado vence en 3 días -> due_soon.count >= 1
        assert body["due_soon"]["count"] >= 1
        assert body["due_soon"]["total"] >= (self._amount_total - self._amount_paid) - 0.01

    def test_g_abono_reduce_saldo(self, vendedor):
        s, _ = vendedor
        eid = TestPayables._entry_id
        # abono 300 sobre saldo 800 -> saldo 500
        r = s.post(f"{BASE_URL}/api/finance/payables/{eid}/payment",
                   params={"amount": 300, "method": "cash", "note": "QA abono"}, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        assert abs(float(r.json()["new_balance"]) - 500.0) < 0.01

        # summary NO cambia (devengado): sigue siendo el total
        r_sum = s.get(f"{BASE_URL}/api/finance/summary", timeout=30)
        assert r_sum.status_code == 200

    def test_g_abono_excede_saldo_400(self, vendedor):
        s, _ = vendedor
        eid = TestPayables._entry_id
        r = s.post(f"{BASE_URL}/api/finance/payables/{eid}/payment",
                   params={"amount": 99999, "method": "cash"}, timeout=30)
        assert r.status_code == 400, f"expected 400 got {r.status_code}: {r.text[:200]}"

    def test_zzz_cleanup_payable(self, vendedor):
        s, _ = vendedor
        eid = TestPayables._entry_id
        if eid:
            try:
                s.delete(f"{BASE_URL}/api/finance/{eid}", timeout=30)
            except Exception:
                pass


# ================= FINANZAS D =================
class TestPurchasesXlsx:
    def test_d_xlsx_report(self, vendedor, qa_supplier):
        s, _ = vendedor
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        # crear un egreso con proveedor para asegurar contenido
        payload = {"type": "egreso", "category": "suppliers", "amount": 50,
                   "description": "QA xlsx", "supplier_id": qa_supplier, "date": today}
        r0 = s.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r0.status_code == 200
        eid = r0.json()["_id"]

        month_start = datetime.now(timezone.utc).strftime("%Y-%m-01")
        r = s.get(f"{BASE_URL}/api/finance/purchases-report.xlsx",
                  params={"date_from": month_start, "date_to": today}, timeout=60)
        assert r.status_code == 200, r.text[:200]
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct, f"bad content-type: {ct}"
        # validate xlsx opens with openpyxl
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert "Por Proveedor" in wb.sheetnames and "Detalle" in wb.sheetnames

        # cleanup
        try:
            s.delete(f"{BASE_URL}/api/finance/{eid}", timeout=30)
        except Exception:
            pass


# ================= CAJA 1, 2, 3 =================
class TestCaja:
    _reg_id = None
    _egreso_id = None
    _sale_id = None

    def test_1_open(self, vendedor):
        s, user = vendedor
        assert user.get("branch_id"), "vendedor no tiene branch_id"
        r = s.post(f"{BASE_URL}/api/cash-register/open",
                   json={"opening_amount": 100, "opening_notes": "QA iter8"}, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        TestCaja._reg_id = r.json()["register"]["_id"]

    def test_1_no_double_open(self, vendedor):
        s, _ = vendedor
        r = s.post(f"{BASE_URL}/api/cash-register/open",
                   json={"opening_amount": 50}, timeout=30)
        assert r.status_code == 400, f"expected 400 got {r.status_code}"

    def test_1_current(self, vendedor):
        s, _ = vendedor
        r = s.get(f"{BASE_URL}/api/cash-register/current", timeout=30)
        assert r.status_code == 200
        assert r.json().get("register", {}).get("_id") == TestCaja._reg_id

    def test_3_egreso_efectivo_durante_caja(self, vendedor, qa_supplier):
        """OBSERVACIONAL: crear un egreso cash durante caja abierta."""
        s, _ = vendedor
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        payload = {"type": "egreso", "category": "suppliers", "amount": 75,
                   "description": "QA egreso durante caja", "supplier_id": qa_supplier,
                   "date": today, "reference": "cash"}
        r = s.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r.status_code == 200
        TestCaja._egreso_id = r.json()["_id"]

    def test_3_preview_no_incluye_egreso(self, vendedor):
        """OBSERVACIONAL CRÍTICO: verificar si el egreso aparece en el preview."""
        s, _ = vendedor
        r = s.get(f"{BASE_URL}/api/cash-register/current/preview", timeout=30)
        assert r.status_code == 200, r.text[:200]
        body = r.json()
        # observamos: no hay campo relativo a egresos ni resta al expected_cash
        print("=" * 60)
        print("[CAJA 3] Preview keys:", sorted(body.keys()))
        print("[CAJA 3] expected_cash:", body.get("expected_cash"),
              " opening=100 (deberia ser 100 si egreso NO se refleja, <100 si SI)")
        print("[CAJA 3] totals_by_method:", body.get("totals_by_method"))
        print("[CAJA 3] payments_detail count:", len(body.get("payments_detail") or []))
        print("[CAJA 3] keys con 'egreso'/'expense':",
              [k for k in body.keys() if "egres" in k.lower() or "expen" in k.lower()])
        print("=" * 60)
        # Assertion documentando comportamiento actual: egreso NO se resta
        # (esperamos expected_cash == opening_amount = 100)
        assert body.get("expected_cash") == 100.0, \
            f"[OBSERVACION] expected_cash != opening: {body.get('expected_cash')} (posible cambio de comportamiento)"

    def test_9_close(self, vendedor):
        s, _ = vendedor
        r = s.post(f"{BASE_URL}/api/cash-register/close",
                   json={"counted_cash": 100, "closing_notes": "QA close"}, timeout=30)
        assert r.status_code == 200, f"{r.status_code}: {r.text[:200]}"
        body = r.json()["register"]
        print("=" * 60)
        print("[CAJA 3 close] keys:", sorted(body.keys()))
        print("[CAJA 3 close] expected_cash:", body.get("expected_cash"),
              "counted_cash:", body.get("counted_cash"),
              "cash_difference:", body.get("cash_difference"))
        print("[CAJA 3 close] totals_by_method:", body.get("totals_by_method"))
        print("[CAJA 3 close] keys con egreso/expense:",
              [k for k in body.keys() if "egres" in k.lower() or "expen" in k.lower()])
        print("=" * 60)
        # OBSERVACION: cash_difference == 0 confirma que egreso NO reduce expected
        assert body.get("expected_cash") == 100.0
        assert body.get("cash_difference") == 0.0, \
            f"cash_difference={body.get('cash_difference')} — si egresos se reflejaran, seria negativo"

    def test_z_cleanup_egreso(self, vendedor):
        s, _ = vendedor
        eid = TestCaja._egreso_id
        if eid:
            try:
                s.delete(f"{BASE_URL}/api/finance/{eid}", timeout=30)
            except Exception:
                pass

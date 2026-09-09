"""Iter14: Estado de Resultados mensual (P&L) - base caja
- GET /api/finance/income-statement (JSON)
- GET /api/finance/income-statement.xlsx (PK, spreadsheet)
- GET /api/finance/income-statement.pdf (%PDF, application/pdf)
- Base caja: ingresos por fecha; egresos contado por fecha; egresos credito SOLO abonos en rango
- Excluye is_voided; superadmin -> 403; branch_id invalido -> 400
- Coexistencia con /api/finance/summary (devengado)
"""
import io
import os
import pytest
import requests
from datetime import datetime, timezone

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ADMIN = ("admin@cortexia.gt", "DemoAdmin2026!")
VEND = ("vendedor@cortexia.gt", "DemoUser2026!")
SUPER = ("superadmin@cortexia.com", "Montecristo2026")

_state = {"ids": [], "supplier_id": None, "credit_id": None,
          "date": datetime.now(timezone.utc).strftime("%Y-%m-%d")}


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text[:200]}"
    csrf = s.cookies.get("csrf_token")
    if csrf:
        s.headers.update({"X-CSRF-Token": csrf})
    return s


@pytest.fixture(scope="module")
def admin():
    return _login(*ADMIN)


@pytest.fixture(scope="module")
def vendedor():
    return _login(*VEND)


@pytest.fixture(scope="module")
def superadmin():
    return _login(*SUPER)


class TestSeed:
    def test_01_create_supplier(self, admin):
        r = admin.post(f"{BASE_URL}/api/suppliers",
                       json={"name": "QA14 Proveedor", "phone": "0000"}, timeout=30)
        assert r.status_code == 200, r.text[:200]
        _state["supplier_id"] = r.json()["_id"]

    def test_02_seed_income_500(self, admin):
        r = admin.post(f"{BASE_URL}/api/finance", json={
            "type": "ingreso", "category": "other_income", "amount": 500,
            "description": "QA14 ingreso otros", "date": _state["date"], "payment_method": "cash",
        }, timeout=30)
        assert r.status_code == 200, r.text[:200]
        _state["ids"].append(r.json()["_id"])

    def test_03_seed_expense_cash_300_rent(self, admin):
        r = admin.post(f"{BASE_URL}/api/finance", json={
            "type": "egreso", "category": "rent", "amount": 300,
            "description": "QA14 alquiler contado", "date": _state["date"], "payment_method": "cash",
        }, timeout=30)
        assert r.status_code == 200, r.text[:200]
        _state["ids"].append(r.json()["_id"])

    def test_04_seed_expense_credit_1000_paid_200(self, admin):
        payload = {
            "type": "egreso", "category": "suppliers", "amount": 1000,
            "description": "QA14 proveedor credito", "date": _state["date"],
            "is_credit": True, "supplier_id": _state["supplier_id"],
            "amount_paid": 200, "payment_method": "cash",
        }
        r = admin.post(f"{BASE_URL}/api/finance", json=payload, timeout=30)
        assert r.status_code == 200, r.text[:200]
        _state["credit_id"] = r.json()["_id"]
        _state["ids"].append(r.json()["_id"])


class TestJSON:
    def test_10_admin_json_shape_and_amounts(self, admin):
        d = _state["date"]
        r = admin.get(f"{BASE_URL}/api/finance/income-statement",
                      params={"date_from": d, "date_to": d}, timeout=30)
        assert r.status_code == 200, r.text[:200]
        j = r.json()
        for k in ("period", "basis", "income", "expenses", "total_income", "total_expense", "net_profit"):
            assert k in j, f"missing key {k}: {j}"
        assert j["basis"] == "caja"
        assert j["period"]["from"] == d and j["period"]["to"] == d
        # DELTA: nuestros egresos QA14 -> rent=300 + suppliers=200 (base caja, credito solo lo abonado)
        exp = {row["category"]: row["amount"] for row in j["expenses"]}
        assert exp.get("rent", 0) >= 300, f"rent expected>=300, got {exp}"
        # suppliers puede ser exactamente 200 (nuestro pago) - si hay otros egresos suppliers en el dia, seria >=200
        assert exp.get("suppliers", 0) >= 200 and exp.get("suppliers", 0) < 1000, \
            f"suppliers cash-basis expected 200 (not 1000 devengado), got {exp}"
        # ingresos: other_income >= 500
        inc = {row["category"]: row["amount"] for row in j["income"]}
        assert inc.get("other_income", 0) >= 500, f"other_income>=500 got {inc}"

    def test_11_vendedor_json_200(self, vendedor):
        d = _state["date"]
        r = vendedor.get(f"{BASE_URL}/api/finance/income-statement",
                        params={"date_from": d, "date_to": d}, timeout=30)
        assert r.status_code == 200

    def test_12_superadmin_403(self, superadmin):
        r = superadmin.get(f"{BASE_URL}/api/finance/income-statement", timeout=30)
        assert r.status_code == 403

    def test_13_branch_invalid_400(self, admin):
        r = admin.get(f"{BASE_URL}/api/finance/income-statement",
                      params={"branch_id": "not-an-oid"}, timeout=30)
        assert r.status_code == 400


class TestXLSX:
    def test_20_admin_xlsx(self, admin):
        d = _state["date"]
        r = admin.get(f"{BASE_URL}/api/finance/income-statement.xlsx",
                      params={"date_from": d, "date_to": d}, timeout=60)
        assert r.status_code == 200, r.text[:200]
        assert r.content[:2] == b"PK"
        assert XLSX_CT in r.headers.get("content-type", "").lower()
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        # Expect at least one sheet with 'Estado' in the title (or the required sheet)
        titles = [ws.title for ws in wb.worksheets]
        assert any("Estado" in t or "Resultados" in t for t in titles), titles
        # Verify labels INGRESOS/EGRESOS/UTILIDAD NETA appear
        flat = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for c in row:
                    if c is not None:
                        flat.append(str(c).upper())
        text = " | ".join(flat)
        assert "INGRESOS" in text
        assert "EGRESOS" in text
        assert "UTILIDAD NETA" in text or "UTILIDAD" in text

    def test_21_superadmin_xlsx_403(self, superadmin):
        r = superadmin.get(f"{BASE_URL}/api/finance/income-statement.xlsx", timeout=30)
        assert r.status_code == 403

    def test_22_xlsx_branch_invalid_400(self, admin):
        r = admin.get(f"{BASE_URL}/api/finance/income-statement.xlsx",
                      params={"branch_id": "not-an-oid"}, timeout=30)
        assert r.status_code == 400


class TestPDF:
    def test_30_admin_pdf(self, admin):
        d = _state["date"]
        r = admin.get(f"{BASE_URL}/api/finance/income-statement.pdf",
                      params={"date_from": d, "date_to": d}, timeout=60)
        assert r.status_code == 200, r.text[:200]
        assert r.content[:4] == b"%PDF"
        assert "application/pdf" in r.headers.get("content-type", "").lower()

    def test_31_superadmin_pdf_403(self, superadmin):
        r = superadmin.get(f"{BASE_URL}/api/finance/income-statement.pdf", timeout=30)
        assert r.status_code == 403

    def test_32_pdf_branch_invalid_400(self, admin):
        r = admin.get(f"{BASE_URL}/api/finance/income-statement.pdf",
                      params={"branch_id": "not-an-oid"}, timeout=30)
        assert r.status_code == 400


class TestAccrualVsCash:
    """/api/finance/summary (devengado) cuenta credito por TOTAL Q1000
    /api/finance/income-statement (caja) cuenta credito solo por pagado Q200
    """
    def test_40_summary_devengado_includes_full_credit(self, admin):
        d = _state["date"]
        rs = admin.get(f"{BASE_URL}/api/finance/summary",
                       params={"date_from": d, "date_to": d}, timeout=30)
        assert rs.status_code == 200, rs.text[:200]
        summary = rs.json()
        # Get income-statement (caja) para comparar
        ri = admin.get(f"{BASE_URL}/api/finance/income-statement",
                       params={"date_from": d, "date_to": d}, timeout=30)
        assert ri.status_code == 200
        pl = ri.json()
        # devengado.total_expense >= caja.total_expense + 800 (diff del credito no pagado)
        # devengado cuenta 1000, caja cuenta 200 -> diff 800
        summary_expense = float(summary.get("expense", summary.get("total_expense", 0)) or 0)
        cash_expense = float(pl["total_expense"])
        assert summary_expense - cash_expense >= 799.99, \
            f"expected devengado - caja >= 800 (credito no pagado), got summary={summary_expense}, caja={cash_expense}"


class TestZZCleanup:
    def test_zz(self):
        try:
            import asyncio
            from motor.motor_asyncio import AsyncIOMotorClient
            from bson import ObjectId as _OID
            mongo_url = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
            db_name = os.environ.get("DB_NAME", "opticas_saas")

            async def _clean():
                client = AsyncIOMotorClient(mongo_url)
                db = client[db_name]
                eids = [_OID(x) for x in _state["ids"] if x]
                if eids:
                    await db.finance_entries.delete_many({"_id": {"$in": eids}})
                await db.finance_entries.delete_many({"description": {"$regex": "^QA14"}})
                if _state.get("supplier_id"):
                    try:
                        await db.suppliers.delete_one({"_id": _OID(_state["supplier_id"])})
                    except Exception:
                        pass
                await db.suppliers.delete_many({"name": {"$regex": "^QA14"}})
                client.close()

            try:
                loop = asyncio.get_event_loop()
                if loop.is_closed():
                    raise RuntimeError("closed")
            except RuntimeError:
                loop = asyncio.new_event_loop()
                asyncio.set_event_loop(loop)
            loop.run_until_complete(_clean())
        except Exception as ex:
            print(f"cleanup err: {ex}")

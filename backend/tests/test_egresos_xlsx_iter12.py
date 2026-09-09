"""Iter12: GET /api/cash-register/{register_id}/egresos.xlsx
- 200 xlsx (PK magic) para caja ABIERTA (recalcula en vivo) y CERRADA (snapshot)
- superadmin 403, id invalido 400, id de otra empresa 404
- El .xlsx contiene descripciones y filas TOTAL/En efectivo (validado con openpyxl si disponible)
- Egresos anulados NO aparecen
"""
import io
import os
import pytest
import requests
from datetime import datetime, timezone

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ADMIN = ("admin@cortexia.gt", "DemoAdmin2026!")
VEND = ("vendedor@cortexia.gt", "DemoUser2026!")
SUPER = ("superadmin@cortexia.com", "Montecristo2026")


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text[:200]}"
    csrf = s.cookies.get("csrf_token")
    if csrf:
        s.headers.update({"X-CSRF-Token": csrf})
    return s


def _ensure_closed(sess):
    try:
        r = sess.get(f"{BASE_URL}/api/cash-register/current", timeout=30)
        if r.status_code == 200 and r.json().get("register"):
            sess.post(f"{BASE_URL}/api/cash-register/close", json={"counted_cash": 0, "closing_notes": "cleanup iter12"}, timeout=30)
    except Exception:
        pass


def _xlsx_text(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(out)
    except Exception as e:
        return f"__xlsx_parse_err:{e}"


_state = {"register_id": None, "egresos_ids": [], "voided_id": None,
          "descs": ["QA12 cash egreso alpha", "QA12 transfer beta", "QA12 anulado gamma"]}


@pytest.fixture(scope="module")
def admin():
    return _login(*ADMIN)


@pytest.fixture(scope="module")
def vendedor():
    s = _login(*VEND)
    _ensure_closed(s)
    yield s
    _ensure_closed(s)


@pytest.fixture(scope="module")
def superadmin():
    return _login(*SUPER)


class TestOpenAndPopulate:
    def test_01_open(self, vendedor):
        r = vendedor.post(f"{BASE_URL}/api/cash-register/open", json={"opening_amount": 100, "opening_notes": "QA12"}, timeout=30)
        assert r.status_code == 200, r.text[:200]
        _state["register_id"] = r.json()["register"]["_id"]

    def test_02_egresos(self, vendedor):
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        for amt, method, desc in [(30, "cash", _state["descs"][0]),
                                  (20, "transfer", _state["descs"][1]),
                                  (10, "cash", _state["descs"][2])]:
            r = vendedor.post(f"{BASE_URL}/api/finance", json={
                "type": "egreso", "category": "other_expense", "amount": amt,
                "description": desc, "date": today, "payment_method": method,
            }, timeout=30)
            assert r.status_code == 200, r.text[:200]
            _state["egresos_ids"].append(r.json()["_id"])
        # anular el tercero
        vid = _state["egresos_ids"][2]
        _state["voided_id"] = vid
        r = vendedor.post(f"{BASE_URL}/api/finance/{vid}/void", json={"reason": "QA12 anular"}, timeout=30)
        assert r.status_code == 200, r.text[:200]


class TestOpenXlsx:
    def test_03_xlsx_open_admin(self, admin):
        rid = _state["register_id"]
        r = admin.get(f"{BASE_URL}/api/cash-register/{rid}/egresos.xlsx", timeout=60)
        assert r.status_code == 200, r.text[:200]
        assert r.content[:2] == b"PK", "not xlsx PK"
        assert XLSX_CT in r.headers.get("content-type", "").lower()
        text = _xlsx_text(r.content)
        assert _state["descs"][0][:15] in text, text[:800]
        assert _state["descs"][1][:15] in text, text[:800]
        # anulado NO aparece
        assert _state["descs"][2][:15] not in text, f"anulado presente!\n{text[:800]}"
        # totales: TOTAL 50 y En efectivo 30 (30 cash + 20 transfer, anulado 10 excluido)
        assert "TOTAL" in text.upper()
        assert "efectivo" in text.lower() or "Efectivo" in text
        # numeric totals
        assert "50" in text
        assert "30" in text


class TestPermissions:
    def test_04_superadmin_403(self, superadmin):
        rid = _state["register_id"]
        r = superadmin.get(f"{BASE_URL}/api/cash-register/{rid}/egresos.xlsx", timeout=30)
        assert r.status_code == 403, f"expected 403 got {r.status_code}"

    def test_05_invalid_id_400(self, admin):
        r = admin.get(f"{BASE_URL}/api/cash-register/not-an-oid/egresos.xlsx", timeout=30)
        assert r.status_code == 400, r.text[:200]

    def test_06_other_company_404(self, admin):
        r = admin.get(f"{BASE_URL}/api/cash-register/507f1f77bcf86cd799439011/egresos.xlsx", timeout=30)
        assert r.status_code == 404, r.text[:200]


class TestClosedXlsx:
    def test_07_close(self, vendedor):
        r = vendedor.post(f"{BASE_URL}/api/cash-register/close", json={"counted_cash": 70, "closing_notes": "QA12"}, timeout=30)
        assert r.status_code == 200, r.text[:200]

    def test_08_xlsx_closed(self, admin):
        rid = _state["register_id"]
        r = admin.get(f"{BASE_URL}/api/cash-register/{rid}/egresos.xlsx", timeout=60)
        assert r.status_code == 200
        assert r.content[:2] == b"PK"
        text = _xlsx_text(r.content)
        assert _state["descs"][0][:15] in text
        assert _state["descs"][1][:15] in text
        assert _state["descs"][2][:15] not in text


class TestZZCleanup:
    def test_zz(self):
        try:
            import asyncio
            from motor.motor_asyncio import AsyncIOMotorClient
            from bson import ObjectId as _OID
            mongo_url = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
            db_name = os.environ.get("DB_NAME", "opticas_saas")

            async def _clean():
                client = AsyncIOMotorClient(mongo_url)
                db = client[db_name]
                eids = [_OID(x) for x in _state["egresos_ids"]]
                if eids:
                    await db.finance_entries.delete_many({"_id": {"$in": eids}})
                if _state.get("register_id"):
                    await db.cash_registers.delete_one({"_id": _OID(_state["register_id"])})
                client.close()

            try:
                loop = asyncio.get_event_loop()
                if loop.is_closed():
                    raise RuntimeError("closed")
            except RuntimeError:
                loop = asyncio.new_event_loop()
                asyncio.set_event_loop(loop)
            loop.run_until_complete(_clean())
        except Exception as ex:
            print(f"cleanup err: {ex}")
